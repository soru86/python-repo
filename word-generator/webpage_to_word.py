#!/usr/bin/env python3
"""
Webpage to Word Document Converter

This script converts a webpage URL to a Microsoft Word document (.docx)
while preserving styling, formatting, tables, fonts, and other elements.
"""

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re
import sys
import os
from urllib.parse import urljoin, urlparse
import http.cookiejar


class WebpageToWordConverter:
    """Converts webpage HTML to Word document with styling preservation."""
    
    def __init__(self, url, cookies=None, auth=None):
        self.url = url
        self.doc = Document()
        self.base_url = url
        self.session = requests.Session()
        
        # Set comprehensive browser-like headers to avoid 403 errors
        self.session.headers.update({
            'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8,application/signed-exchange;v=b3;q=0.7',
            'Accept-Language': 'en-US,en;q=0.9',
            'Accept-Encoding': 'gzip, deflate, br',
            'DNT': '1',
            'Connection': 'keep-alive',
            'Upgrade-Insecure-Requests': '1',
            'Sec-Fetch-Dest': 'document',
            'Sec-Fetch-Mode': 'navigate',
            'Sec-Fetch-Site': 'none',
            'Sec-Fetch-User': '?1',
            'Cache-Control': 'max-age=0'
        })
        
        # Add cookies if provided
        if cookies:
            if isinstance(cookies, str):
                # If cookies is a string, try to parse it
                if os.path.isfile(cookies):
                    # Load cookies from file
                    jar = http.cookiejar.MozillaCookieJar(cookies)
                    jar.load(ignore_discard=True, ignore_expires=True)
                    self.session.cookies.update(jar)
                else:
                    # Parse cookie string (format: "name1=value1; name2=value2")
                    for cookie in cookies.split(';'):
                        if '=' in cookie:
                            name, value = cookie.strip().split('=', 1)
                            self.session.cookies.set(name, value)
            elif isinstance(cookies, dict):
                self.session.cookies.update(cookies)
        
        # Add authentication if provided
        if auth:
            if isinstance(auth, tuple) and len(auth) == 2:
                self.session.auth = auth
            elif isinstance(auth, dict) and 'username' in auth and 'password' in auth:
                self.session.auth = (auth['username'], auth['password'])
    
    def fetch_webpage(self):
        """Fetch the webpage content."""
        try:
            # Add Referer header based on the URL domain
            parsed_url = urlparse(self.url)
            referer = f"{parsed_url.scheme}://{parsed_url.netloc}/"
            headers = {'Referer': referer}
            
            response = self.session.get(self.url, headers=headers, timeout=30, allow_redirects=True)
            
            # Handle 403 Forbidden errors with helpful suggestions
            if response.status_code == 403:
                print(f"\n❌ Error: 403 Forbidden - Access denied to {self.url}")
                print("\nPossible solutions:")
                print("1. The page requires authentication:")
                print("   - Export cookies from your browser and use them:")
                print("     python webpage_to_word.py <URL> --cookies 'cookie1=value1; cookie2=value2'")
                print("   - Or use browser extension to export cookies to a file")
                print("\n2. The page may require login:")
                print("   - Log in to the website in your browser first")
                print("   - Export cookies and use them with this script")
                print("\n3. For Confluence pages:")
                print("   - Make sure you're logged in to Confluence")
                print("   - Export cookies from your browser session")
                print("   - Use: python webpage_to_word.py <URL> --cookies-file cookies.txt")
                print("\n4. The server may be blocking automated requests:")
                print("   - Try accessing the page in a browser first")
                print("   - Check if the page is publicly accessible")
                sys.exit(1)
            
            response.raise_for_status()
            response.encoding = response.apparent_encoding
            return response.text
        except requests.exceptions.HTTPError as e:
            if e.response.status_code == 403:
                print(f"\n❌ Error: 403 Forbidden - Access denied")
                print(f"URL: {self.url}")
                print("\nThis page requires authentication or has restricted access.")
                print("Please see the suggestions above for resolving this issue.")
            else:
                print(f"❌ HTTP Error {e.response.status_code}: {e}")
            sys.exit(1)
        except requests.exceptions.RequestException as e:
            print(f"❌ Error fetching webpage: {e}")
            print(f"URL: {self.url}")
            sys.exit(1)
    
    def parse_html(self, html_content):
        """Parse HTML content using BeautifulSoup."""
        return BeautifulSoup(html_content, 'html.parser')
    
    def get_font_size(self, element):
        """Extract font size from element style or attributes."""
        style = element.get('style', '')
        if 'font-size' in style:
            match = re.search(r'font-size:\s*(\d+(?:\.\d+)?)(px|pt|em)', style)
            if match:
                size = float(match.group(1))
                unit = match.group(2)
                if unit == 'px':
                    return Pt(size * 0.75)  # Convert px to pt (approximate)
                elif unit == 'pt':
                    return Pt(size)
                elif unit == 'em':
                    return Pt(size * 12)  # Assume base font size of 12pt
        return None
    
    def get_font_family(self, element):
        """Extract font family from element style."""
        style = element.get('style', '')
        if 'font-family' in style:
            match = re.search(r'font-family:\s*([^;]+)', style)
            if match:
                return match.group(1).strip().strip('"\'')
        return None
    
    def get_color(self, element):
        """Extract text color from element style."""
        style = element.get('style', '')
        if 'color' in style:
            match = re.search(r'color:\s*([^;]+)', style)
            if match:
                color_str = match.group(1).strip()
                # Handle hex colors
                if color_str.startswith('#'):
                    hex_color = color_str[1:]
                    if len(hex_color) == 6:
                        r = int(hex_color[0:2], 16)
                        g = int(hex_color[2:4], 16)
                        b = int(hex_color[4:6], 16)
                        return RGBColor(r, g, b)
                # Handle rgb() colors
                rgb_match = re.search(r'rgb\((\d+),\s*(\d+),\s*(\d+)\)', color_str)
                if rgb_match:
                    return RGBColor(int(rgb_match.group(1)), 
                                   int(rgb_match.group(2)), 
                                   int(rgb_match.group(3)))
        return None
    
    def get_alignment(self, element):
        """Extract text alignment from element style or attributes."""
        style = element.get('style', '')
        align_attr = element.get('align', '').lower()
        
        if 'text-align' in style:
            match = re.search(r'text-align:\s*(\w+)', style)
            if match:
                align = match.group(1).lower()
            else:
                align = align_attr
        else:
            align = align_attr
        
        alignment_map = {
            'left': WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
            'justify': WD_ALIGN_PARAGRAPH.JUSTIFY
        }
        return alignment_map.get(align, None)
    
    def apply_formatting(self, run, element):
        """Apply formatting to a text run based on element styles."""
        # Bold
        if element.name in ['b', 'strong'] or element.get('style', '').find('font-weight: bold') != -1:
            run.bold = True
        
        # Italic
        if element.name in ['i', 'em'] or element.get('style', '').find('font-style: italic') != -1:
            run.italic = True
        
        # Underline
        if element.name == 'u' or element.get('style', '').find('text-decoration: underline') != -1:
            run.underline = True
        
        # Font size
        font_size = self.get_font_size(element)
        if font_size:
            run.font.size = font_size
        
        # Font family
        font_family = self.get_font_family(element)
        if font_family:
            run.font.name = font_family.split(',')[0].strip()
        
        # Color
        color = self.get_color(element)
        if color:
            run.font.color.rgb = color
    
    def process_text(self, element, paragraph=None):
        """Process text content from an element."""
        if paragraph is None:
            paragraph = self.doc.add_paragraph()
        
        # Apply paragraph-level formatting
        alignment = self.get_alignment(element)
        if alignment:
            paragraph.alignment = alignment
        
        # Process text nodes and child elements
        for content in element.children:
            if isinstance(content, str):
                text = content.strip()
                if text:
                    run = paragraph.add_run(text)
                    self.apply_formatting(run, element)
            elif hasattr(content, 'name'):
                if content.name in ['br']:
                    paragraph.add_run().add_break()
                elif content.name in ['b', 'strong', 'i', 'em', 'u', 'span', 'a', 'code']:
                    run = paragraph.add_run(content.get_text())
                    self.apply_formatting(run, content)
                elif content.name in ['p', 'div', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'li']:
                    # Recursive processing for block elements
                    self.process_element(content)
                else:
                    # For other elements, just get the text
                    text = content.get_text()
                    if text.strip():
                        run = paragraph.add_run(text)
                        self.apply_formatting(run, element)
        
        return paragraph
    
    def process_heading(self, element):
        """Process heading elements (h1-h6)."""
        level_map = {
            'h1': 1, 'h2': 2, 'h3': 3,
            'h4': 4, 'h5': 5, 'h6': 6
        }
        level = level_map.get(element.name, 1)
        heading = self.doc.add_heading(level=level)
        
        # Process text with formatting
        for content in element.children:
            if isinstance(content, str):
                text = content.strip()
                if text:
                    run = heading.add_run(text)
                    self.apply_formatting(run, element)
            elif hasattr(content, 'name'):
                run = heading.add_run(content.get_text())
                self.apply_formatting(run, content)
        
        # Apply alignment
        alignment = self.get_alignment(element)
        if alignment:
            heading.alignment = alignment
    
    def process_table(self, table_element):
        """Process HTML table and convert to Word table."""
        # Find all rows
        rows = table_element.find_all(['tr'], recursive=False)
        if not rows:
            return
        
        # Determine number of columns
        max_cols = 0
        for row in rows:
            cols = len(row.find_all(['td', 'th'], recursive=False))
            max_cols = max(max_cols, cols)
        
        if max_cols == 0:
            return
        
        # Create Word table
        word_table = self.doc.add_table(rows=len(rows), cols=max_cols)
        word_table.style = 'Light Grid Accent 1'
        
        # Populate table
        for row_idx, row in enumerate(rows):
            cells = row.find_all(['td', 'th'], recursive=False)
            for col_idx, cell in enumerate(cells):
                if col_idx < max_cols:
                    word_cell = word_table.rows[row_idx].cells[col_idx]
                    word_cell.text = cell.get_text(strip=True)
                    
                    # Apply cell formatting
                    if cell.name == 'th':
                        for paragraph in word_cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                    
                    # Apply alignment
                    alignment = self.get_alignment(cell)
                    if alignment:
                        word_cell.paragraphs[0].alignment = alignment
                    
                    # Apply color
                    color = self.get_color(cell)
                    if color:
                        for paragraph in word_cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.color.rgb = color
    
    def process_list(self, list_element):
        """Process HTML lists (ul/ol) and convert to Word lists."""
        is_ordered = list_element.name == 'ol'
        items = list_element.find_all('li', recursive=False)
        
        for item in items:
            paragraph = self.doc.add_paragraph(style='List Bullet' if not is_ordered else 'List Number')
            self.process_text(item, paragraph)
    
    def process_element(self, element):
        """Process a single HTML element and convert to Word document elements."""
        if element.name is None:
            return
        
        tag_name = element.name.lower()
        
        if tag_name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            self.process_heading(element)
        elif tag_name == 'table':
            self.process_table(element)
        elif tag_name in ['ul', 'ol']:
            self.process_list(element)
        elif tag_name in ['p', 'div', 'section', 'article', 'main', 'header', 'footer']:
            self.process_text(element)
        elif tag_name in ['br']:
            self.doc.add_paragraph()
        elif tag_name in ['hr']:
            paragraph = self.doc.add_paragraph()
            paragraph.add_run('_' * 50)
        elif tag_name in ['script', 'style', 'meta', 'link', 'noscript']:
            # Skip these elements
            pass
        else:
            # For other elements, try to process as text
            text = element.get_text(strip=True)
            if text:
                self.process_text(element)
    
    def convert(self):
        """Main conversion method."""
        print(f"Fetching webpage: {self.url}")
        html_content = self.fetch_webpage()
        
        print("Parsing HTML content...")
        soup = self.parse_html(html_content)
        
        # Remove script and style elements
        for script in soup(["script", "style"]):
            script.decompose()
        
        # Get the main content
        body = soup.find('body')
        if body:
            print("Converting HTML to Word document...")
            for element in body.children:
                if hasattr(element, 'name') and element.name:
                    self.process_element(element)
        else:
            # If no body tag, process the whole document
            print("Converting HTML to Word document...")
            for element in soup.children:
                if hasattr(element, 'name') and element.name:
                    self.process_element(element)
        
        return self.doc
    
    def save(self, filename=None):
        """Save the Word document."""
        if filename is None:
            # Generate filename from URL
            parsed_url = urlparse(self.url)
            domain = parsed_url.netloc.replace('www.', '')
            path = parsed_url.path.strip('/').replace('/', '_')
            if not path:
                path = 'index'
            filename = f"{domain}_{path}.docx"
            # Clean filename
            filename = re.sub(r'[^\w\-_\.]', '_', filename)
            if len(filename) > 200:
                filename = filename[:200] + '.docx'
        
        # Ensure .docx extension
        if not filename.endswith('.docx'):
            filename += '.docx'
        
        filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)
        self.doc.save(filepath)
        print(f"Word document saved: {filepath}")
        return filepath


def main():
    """Main function to run the script."""
    import argparse
    
    parser = argparse.ArgumentParser(
        description='Convert a webpage URL to a Microsoft Word document (.docx)',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Basic usage
  python webpage_to_word.py https://example.com
  
  # With cookies (for authenticated pages)
  python webpage_to_word.py https://confluence.example.com/page --cookies "JSESSIONID=abc123; atlassian.xsrf.token=xyz"
  
  # With cookies file (exported from browser)
  python webpage_to_word.py https://confluence.example.com/page --cookies-file cookies.txt
  
  # With basic authentication
  python webpage_to_word.py https://example.com --username user --password pass

Note: For Confluence or other authenticated pages, you'll need to export cookies
from your browser. Use browser extensions like "Cookie-Editor" or "EditThisCookie"
to export cookies, then use --cookies or --cookies-file option.
        """
    )
    
    parser.add_argument('url', help='URL of the webpage to convert')
    parser.add_argument('--cookies', '-c', help='Cookies string (format: "name1=value1; name2=value2")')
    parser.add_argument('--cookies-file', '-f', help='Path to cookies file (Netscape/Mozilla format)')
    parser.add_argument('--username', '-u', help='Username for basic authentication')
    parser.add_argument('--password', '-p', help='Password for basic authentication')
    
    args = parser.parse_args()
    
    url = args.url
    
    # Validate URL
    if not url.startswith(('http://', 'https://')):
        url = 'https://' + url
    
    # Prepare cookies
    cookies = None
    if args.cookies:
        cookies = args.cookies
    elif args.cookies_file:
        cookies = args.cookies_file
    
    # Prepare authentication
    auth = None
    if args.username and args.password:
        auth = (args.username, args.password)
    elif args.username:
        import getpass
        password = getpass.getpass(f"Password for {args.username}: ")
        auth = (args.username, password)
    
    converter = WebpageToWordConverter(url, cookies=cookies, auth=auth)
    converter.convert()
    converter.save()


if __name__ == '__main__':
    main()

